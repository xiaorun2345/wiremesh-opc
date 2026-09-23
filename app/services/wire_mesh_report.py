from __future__ import annotations

"""Generate customer-facing wire-mesh quotation PDF and Excel files."""
import io
import uuid
from datetime import datetime



PRODUCT_LABELS = {
    "welded_mesh": "焊接网", "fence_mesh": "护栏网", "hexagonal_mesh": "六角网",
    "screen_mesh": "工业筛网", "expanded_mesh": "钢板网",
}
MATERIAL_LABELS = {
    "low_carbon_steel": "低碳钢丝", "galvanized_steel": "镀锌钢丝",
    "stainless_304": "不锈钢 304", "stainless_316": "不锈钢 316",
}


def _build_excel(quote: dict, task_id: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.merge_cells("A1:E1")
    ws["A1"] = "丝网产品报价单 / WIRE MESH QUOTATION"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="163A63")
    ws["A1"].alignment = Alignment(horizontal="center")
    rows = [
        ("报价编号", task_id), ("生成日期", datetime.now().strftime("%Y-%m-%d")),
        ("产品", PRODUCT_LABELS.get(quote["spec"]["product_type"], quote["spec"]["product_type"])),
        ("材料", MATERIAL_LABELS.get(quote["spec"]["material"], quote["spec"]["material"])),
        ("丝径", f"{quote['spec']['wire_diameter_mm']} mm"),
        ("网孔", f"{quote['spec']['mesh_size_mm']} mm"),
        ("尺寸", f"{quote['spec']['width_m']} × {quote['spec']['length_m']} m"),
        ("数量", quote["spec"]["quantity"]), ("总面积", f"{quote['area_m2']} m²"),
        ("理论重量", f"{quote['total_weight_kg']} kg"),
        ("单价", f"{quote['currency']} {quote['unit_price_m2']} / m²"),
        ("报价总额", f"{quote['currency']} {quote['quote_total']}"),
        ("交期", "以工厂确认结果为准"), ("报价有效期", "7 天"),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        ws.cell(index, 1, label).font = Font(bold=True)
        ws.cell(index, 2, value)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 42
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _build_pdf(quote: dict, task_id: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm, topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("title-cn", parent=styles["Title"], fontName="STSong-Light", fontSize=18, textColor=colors.HexColor("#163A63"), alignment=TA_CENTER)
    body = ParagraphStyle("body-cn", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9, leading=14)
    spec = quote["spec"]
    data = [
        ["项目", "规格 / 内容"],
        ["产品", PRODUCT_LABELS.get(spec["product_type"], spec["product_type"])],
        ["材料", MATERIAL_LABELS.get(spec["material"], spec["material"])],
        ["丝径 / 网孔", f"{spec['wire_diameter_mm']} mm / {spec['mesh_size_mm']} mm"],
        ["尺寸 / 数量", f"{spec['width_m']} × {spec['length_m']} m / {spec['quantity']}"],
        ["总面积 / 重量", f"{quote['area_m2']} m² / {quote['total_weight_kg']} kg"],
        ["表面处理 / 包装", f"{spec['surface_treatment']} / {spec['packaging']}"],
        ["单价", f"{quote['currency']} {quote['unit_price_m2']} / m²"],
        ["报价总额", f"{quote['currency']} {quote['quote_total']}"],
    ]
    table = Table(data, colWidths=[45*mm, 115*mm])
    table.setStyle(TableStyle([
        ("FONTNAME", (0,0), (-1,-1), "STSong-Light"), ("FONTSIZE", (0,0), (-1,-1), 9),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#163A63")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), .5, colors.HexColor("#D8E1EA")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F5F8FB")]),
        ("LEFTPADDING", (0,0), (-1,-1), 8), ("RIGHTPADDING", (0,0), (-1,-1), 8),
        ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ]))
    story = [
        Paragraph("丝网产品报价单", title),
        Paragraph("WIRE MESH QUOTATION", title), Spacer(1, 7*mm),
        Paragraph(f"报价编号：{task_id}　生成日期：{datetime.now():%Y-%m-%d}", body), Spacer(1, 4*mm), table,
        Spacer(1, 6*mm), Paragraph("条款：报价有效期 7 天；交期、付款方式和贸易条款以双方最终合同为准。", body),
        Paragraph("本报价由规则引擎生成并经人工审核后生效。", body),
    ]
    doc.build(story)
    return output.getvalue()


def build_quote_excel(quote: dict, quote_id: str) -> bytes:
    return _build_excel(quote, quote_id)


def build_quote_pdf(quote: dict, quote_id: str) -> bytes:
    return _build_pdf(quote, quote_id)


async def export_wire_mesh_quote(trace_id: str) -> dict:
    from sqlalchemy import select
    from app.core.memory.redis_memory import TaskMemory
    from app.models.report import Report
    from app.services.minio_service import MinioService
    mem = TaskMemory(trace_id)
    quote = await mem.get("wire_mesh_quote")
    if not quote:
        raise ValueError("缺少丝网报价结果")
    excel_data = _build_excel(quote, trace_id)
    pdf_data = _build_pdf(quote, trace_id)
    excel_path = f"reports/{trace_id}/wire-mesh-quotation.xlsx"
    pdf_path = f"reports/{trace_id}/wire-mesh-quotation.pdf"
    storage = MinioService()
    storage.upload(excel_path, excel_data, "wire-mesh-quotation.xlsx")
    storage.upload(pdf_path, pdf_data, "wire-mesh-quotation.pdf")

    import app.utils.db as db_module
    async with db_module.AsyncSessionLocal() as db:
        result = await db.execute(select(Report).where(Report.task_id == trace_id))
        report = result.scalar_one_or_none()
        if report:
            report.pdf_path = pdf_path
            report.excel_path = excel_path
        else:
            db.add(Report(id=str(uuid.uuid4()), task_id=trace_id, pdf_path=pdf_path, excel_path=excel_path))
        await db.commit()
    return {"pdf_path": pdf_path, "excel_path": excel_path, "quote_total": quote["quote_total"]}
